import asyncio
import copy
import io
import json
import os
import subprocess
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError

from app.allocator import allocate
from app.billing import calculate
from app.db import SessionLocal
from app.jobs import process_one
from app.main import app
from app.models import AllocationRun, LoginSession, PlanState, User, now
from app.planning_api import events
from app.planning_rules import NIGHTS, validate
from app.planning_store import snapshot
from app.security import COOKIE, digest
from app.spreadsheet_api import HEADERS

PASSWORD='Synthetic test passphrase 2026!'


def headers(owner):return {'X-CSRF-Token':owner['csrf_token']}
def plan(client):return client.get('/api/planning').json()
def revision(client):return plan(client)['revision']


def member(name='Synthetic Guest',age=30,**fields):
    return {'name':name,'age':age,'category':'guest','arrival':NIGHTS[0],'tags':[],**fields}


def family(client,owner,people=None,**fields):
    body={'revision':revision(client),'label':'Synthetic Family','side':'bride','members':people or [member()],**fields}
    response=client.post('/api/planning/families',json=body,headers=headers(owner));assert response.status_code==201,response.text
    fid=response.json()['family_id']
    return [g for g in plan(client)['guests'] if g['family_id']==fid]


def reserve(client,owner,ids=None):
    response=client.put('/api/planning/reservations',json={'revision':revision(client),'room_ids':ids or ['w-101','w-103','w-201','w-203']},headers=headers(owner));assert response.status_code==200,response.text


def preview(client,owner,guests,rid,night=NIGHTS[0]):
    response=client.post('/api/planning/assignments/preview',json={'revision':revision(client),'moves':[{'guest_id':g['id'],'night':night,'resource_id':rid} for g in guests]},headers=headers(owner));assert response.status_code==200,response.text
    return response.json()


def apply(client,owner,p,ack=True):
    return client.post('/api/planning/assignments/apply',json={'revision':p['revision'],'token':p['token'],'acknowledged':[w['id'] for w in p['warnings']] if ack else []},headers=headers(owner))


def assign(client,owner,guests,rid,night=NIGHTS[0]):
    p=preview(client,owner,guests,rid,night);assert not p['errors'],p
    response=apply(client,owner,p);assert response.status_code==200,response.text


def link(client,owner,a,b,kind='may_share',severity='hard'):
    response=client.post('/api/planning/links',json={'revision':revision(client),'a':a['id'],'b':b['id'],'kind':kind,'severity':severity},headers=headers(owner));assert response.status_code==200,response.text


def test_placeholder_completion_preserves_headcount_and_stale_edit(client,owner):
    people=family(client,owner,[member(None,None) for _ in range(5)])
    d=plan(client);f=d['families'][0]
    members=[{k:g.get(k) for k in ['id','name','age','category','phone','arrival','tags','notes']} for g in people]
    for i,m in enumerate(members):m['name']=f'Synthetic Person {i}';m['age']=20+i
    body={**f,'revision':d['revision'],'members':members};body.pop('id')
    assert client.put('/api/planning/families/'+f['id'],json=body,headers=headers(owner)).status_code==200
    assert len(plan(client)['guests'])==5 and {g['id'] for g in plan(client)['guests']}=={g['id'] for g in people}
    assert client.put('/api/planning/families/'+f['id'],json=body,headers=headers(owner)).status_code==409


def test_concurrent_writes_are_revision_guarded(client,owner):
    rev=revision(client);cookie=client.cookies.get(COOKIE)
    def create(i):
        with TestClient(app) as c:
            c.cookies.set(COOKIE,cookie)
            return c.post('/api/planning/families',json={'revision':rev,'label':f'Synthetic {i}','side':'bride','members':[member()]},headers=headers(owner)).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:assert sorted(pool.map(create,[1,2]))==[201,409]
    assert len(plan(client)['guests'])==1


@pytest.mark.parametrize('role,side',[('family','bride'),('family','groom'),('reception',None)])
def test_scope_no_cost_contact_or_mutation_leak(client,owner,role,side):
    bride=family(client,owner,[member('Synthetic Bride Guest',phone='SYNTHETIC-BRIDE-CONTACT',notes='OWNER-ONLY-NOTE')],side='bride')
    groom=family(client,owner,[member('Synthetic Groom Guest',phone='SYNTHETIC-GROOM-CONTACT')],side='groom',label='Synthetic Groom Family')
    assign(client,owner,bride,'w-102');assign(client,owner,groom,'w-104')
    assert client.post('/api/users',json={'username':'synthetic.viewer','display_name':'Synthetic Viewer','password':PASSWORD,'role':role,'side':side},headers=headers(owner)).status_code==201
    viewer=TestClient(app);auth=viewer.post('/api/auth/login',json={'username':'synthetic.viewer','password':PASSWORD}).json();h={'X-CSRF-Token':auth['csrf_token']}
    response=viewer.get('/api/planning');d=response.json()
    assert 'OWNER-ONLY-NOTE' not in response.text and 'paise' not in response.text
    assert viewer.get('/api/planning/billing').status_code==403
    assert viewer.get('/api/planning/history').status_code==403
    assert viewer.get('/api/planning/allocation').status_code==403
    assert viewer.get('/api/planning/billing/snapshots').status_code==403
    assert viewer.post('/api/planning/families',json={'revision':d['revision'],'label':'Cannot Edit','side':'bride','members':[member()]},headers=h).status_code==403
    assert viewer.put('/api/planning/reservations',json={'revision':d['revision'],'room_ids':['w-101','w-103','w-201','w-203']},headers=h).status_code==403
    assert viewer.post('/api/planning/assignments/preview',json={'revision':d['revision'],'moves':[{'guest_id':bride[0]['id'],'night':NIGHTS[0],'resource_id':'w-104'}]},headers=h).status_code==403
    if role=='family':
        expected=bride if side=='bride' else groom;other=groom if side=='bride' else bride
        assert {g['id'] for g in d['guests']}=={g['id'] for g in expected}
        assert other[0]['name'] not in response.text and other[0]['id'] not in response.text
        assert 'CONTACT' not in response.text and d['summary'][NIGHTS[0]]['free_rooms'] is None
        assert set(d['summary'][NIGHTS[0]]['side_rooms'])=={side}
    else:assert len(d['guests'])==2 and 'SYNTHETIC-BRIDE-CONTACT' in response.text
    book=load_workbook(io.BytesIO(viewer.get('/api/planning/export.xlsx').content))
    assert 'Estimate' not in book.sheetnames
    values=' '.join(str(c.value) for s in book for row in s for c in row)
    assert 'OWNER-ONLY-NOTE' not in values
    if role=='family':assert other[0]['name'] not in values and 'CONTACT' not in values


def test_reservations_capacity_children_bedding_and_ground_floor(client,owner):
    young=family(client,owner,[member('Synthetic Parent A',35),member('Synthetic Parent B',34),member('Synthetic Child A',10),member('Synthetic Child B',4)],couple_children=True)
    assert any(e['code']=='couple_children' for e in preview(client,owner,young,'w-105')['errors'])
    reserve(client,owner)
    assert any(e['code']=='couple_children' for e in preview(client,owner,young,'w-105')['errors'])
    assign(client,owner,young,'w-102')
    s=plan(client)['summary'][NIGHTS[0]];assert s['assigned']==4 and s['extra_bedding']==1
    older=family(client,owner,[member('Synthetic Elder',50)],label='Elder')
    assert any(e['code']=='ground_floor' for e in preview(client,owner,older,'w-202')['errors'])
    assign(client,owner,older,'w-104')
    tagged=family(client,owner,[member('Synthetic Tagged',29,tags=['ground_floor'])],label='Tagged')
    assert any(e['code']=='ground_floor' for e in preview(client,owner,tagged,'w-206')['errors'])
    adults=family(client,owner,[member('Synthetic Adult '+str(i),11+i) for i in range(4)],label='Four adults')
    assert any(e['code']=='capacity' for e in preview(client,owner,adults,'w-110')['errors'])
    unknown=family(client,owner,[member('Unknown Age',None)],label='Unknown')
    assert any(e['code']=='age_unknown' for e in preview(client,owner,unknown,'w-112')['errors'])


def test_dormitory_second_hotel_dates_and_reserved_eligibility(client,owner):
    reserve(client,owner)
    guest=family(client,owner,[member('Synthetic Regular',30)],label='Regular')
    assert any(e['code']=='reserved_family' for e in preview(client,owner,guest,'w-101')['errors'])
    assert any(e['code']=='staff_separation' for e in preview(client,owner,guest,'bed-1')['errors'])
    assert any(e['code']=='second_hotel' for e in preview(client,owner,guest,'second-0-1')['errors'])
    eligible=family(client,owner,[member('Synthetic Eligible',30,tags=['second_hotel'],arrival=NIGHTS[1])],label='Eligible',reserved_eligible=True)
    assert any(e['code']=='stay' for e in preview(client,owner,eligible,'second-0-1')['errors'])
    assign(client,owner,eligible,'second-0-1',NIGHTS[1])
    staff=family(client,owner,[member('Synthetic Driver',30,category='driver')],label='Driver')
    assert any(e['code']=='staff_separation' for e in preview(client,owner,staff,'w-123')['errors'])
    assign(client,owner,staff,'bed-1')
    other=family(client,owner,[member('Synthetic Worker',31,category='staff')],label='Worker')
    assert any(e['code']=='capacity' for e in preview(client,owner,other,'bed-1')['errors'])
    regular_family=family(client,owner,[member('Synthetic Reserved Adult',40)],label='Family reserve',reserved_eligible=True)
    assign(client,owner,regular_family,'w-101')


def test_warning_preview_binds_revision_exact_ack_and_is_single_use(client,owner):
    people=family(client,owner,[member('Synthetic A'),member('Synthetic B')])
    p=preview(client,owner,[people[0]],'w-102');assert p['token'] and any(w['code']=='family_split' for w in p['warnings'])
    assert p['changes']==[{'guest_id':people[0]['id'],'night':NIGHTS[0],'from':None,'to':'w-102'}]
    assert p['estimate']['known_subtotal_paise']==0 and not plan(client)['assignments']
    assert apply(client,owner,p,ack=False).status_code==409
    assert apply(client,owner,p).status_code==200
    assert apply(client,owner,p).status_code==409
    p=preview(client,owner,[people[1]],'w-104')
    family(client,owner,label='Another change')
    assert apply(client,owner,p).status_code==409
    assert len(plan(client)['assignments'])==1


def test_undo_that_restores_a_warning_requires_exact_review(client,owner):
    people=family(client,owner,[member('Synthetic A'),member('Synthetic B')])
    assign(client,owner,[people[0]],'w-102')
    assign(client,owner,[people[1]],'w-102')
    result=client.post('/api/planning/history',json={'revision':revision(client),'direction':'undo'},headers=headers(owner))
    assert result.status_code==200,result.text
    p=result.json();assert p['token'] and p['warnings']
    assert len(plan(client)['assignments'])==2
    assert apply(client,owner,p,ack=False).status_code==409
    assert apply(client,owner,p).status_code==200
    assert len(plan(client)['assignments'])==1
    assert client.post('/api/planning/history',json={'revision':revision(client),'direction':'redo'},headers=headers(owner)).status_code==200
    assert len(plan(client)['assignments'])==2


def test_compatibility_locks_swap_undo_redo_and_candidates(client,owner):
    a=family(client,owner,[member('Synthetic A')],label='A');b=family(client,owner,[member('Synthetic B')],label='B')
    assign(client,owner,a,'w-102')
    assert any(e['code']=='compatibility' for e in preview(client,owner,b,'w-102')['errors'])
    link(client,owner,a[0],b[0]);assign(client,owner,b,'w-104')
    r=client.post('/api/planning/assignments/destinations',json={'revision':revision(client),'guest_ids':[b[0]['id']],'night':NIGHTS[0]},headers=headers(owner)).json()
    assert r['choices']['w-102']['allowed'] is True and r['choices']['bed-1']['allowed'] is False
    lock={'revision':revision(client),'resource_id':'w-102','night':NIGHTS[0],'locked':True}
    assert client.put('/api/planning/locks',json=lock,headers=headers(owner)).status_code==200
    assert any(e['code']=='locked' for e in preview(client,owner,a,'w-110')['errors'])
    lock.update(revision=revision(client),locked=False);assert client.put('/api/planning/locks',json=lock,headers=headers(owner)).status_code==200
    p=client.post('/api/planning/assignments/preview',json={'revision':revision(client),'moves':[{'guest_id':a[0]['id'],'night':NIGHTS[0],'resource_id':'w-104'},{'guest_id':b[0]['id'],'night':NIGHTS[0],'resource_id':'w-102'}]},headers=headers(owner)).json()
    assert apply(client,owner,p).status_code==200
    rows=plan(client)['assignments'];assert next(r for r in rows if r['guest_id']==a[0]['id'])['resource_id']=='w-104'
    for direction,resource in [('undo','w-102'),('redo','w-104')]:
        assert client.post('/api/planning/history',json={'revision':revision(client),'direction':direction},headers=headers(owner)).status_code==200
        assert next(r for r in plan(client)['assignments'] if r['guest_id']==a[0]['id'])['resource_id']==resource


def test_conditional_rules_verified_connections_and_must_share(client,owner):
    a=family(client,owner,[member('Synthetic A')],label='A');b=family(client,owner,[member('Synthetic B')],label='B')
    link(client,owner,a[0],b[0],'nearby')
    p=preview(client,owner,a,'w-102');assert any(e['code'].startswith('link_') for e in p['errors'])
    connection={'revision':revision(client),'a':'w-102','b':'w-104','kind':'opposite','confirmed':True}
    assert client.post('/api/planning/connections',json=connection,headers=headers(owner)).status_code==200
    p=client.post('/api/planning/assignments/preview',json={'revision':revision(client),'moves':[{'guest_id':a[0]['id'],'night':NIGHTS[0],'resource_id':'w-102'},{'guest_id':b[0]['id'],'night':NIGHTS[0],'resource_id':'w-104'}]},headers=headers(owner)).json();assert not p['errors'];assert apply(client,owner,p).status_code==200
    assert client.post('/api/planning/rules',json={'revision':revision(client),'name':'Synthetic allowed block','kind':'allowed_resources','side':'groom','resource_ids':['w-110']},headers=headers(owner)).status_code==200
    groom=family(client,owner,[member('Synthetic Groom')],side='groom',label='Groom')
    assert any(e['code'].startswith('rule_') for e in preview(client,owner,groom,'w-112')['errors'])
    assign(client,owner,groom,'w-110')


def test_worker_proposal_preserves_locks_scope_and_stale_apply(client,owner):
    reserve(client,owner)
    elder=family(client,owner,[member('Synthetic Elder',55)],label='Elder')
    children=family(client,owner,[member('Synthetic Parent A',35),member('Synthetic Parent B',34),member('Synthetic Child A',8),member('Synthetic Child B',10)],label='Parents',couple_children=True)
    staff=family(client,owner,[member('Synthetic Staff',35,category='staff')],label='Staff')
    unknown=family(client,owner,[member('Synthetic Unknown',None)],label='Unknown')
    assign(client,owner,elder,'w-104')
    assert client.put('/api/planning/locks',json={'revision':revision(client),'resource_id':'w-104','night':NIGHTS[0],'locked':True},headers=headers(owner)).status_code==200
    response=client.post('/api/planning/allocation',json={'revision':revision(client)},headers=headers(owner));assert response.status_code==202
    run_id=response.json()['id'];before=plan(client)['assignments'];assert process_one()
    runs=client.get('/api/planning/allocation').json()['runs'];job=next(j for j in runs if j['id']==run_id)
    assert job['status']=='completed' and job['result']['assignments'] is not None,job
    proposal=job['result']['assignments'];assert plan(client)['assignments']==before
    assert any(a['guest_id']==elder[0]['id'] and a['night']==NIGHTS[0] and a['resource_id']=='w-104' for a in proposal)
    assert any(a['guest_id']==staff[0]['id'] and a['resource_id'].startswith('bed-') for a in proposal)
    assert any(u['guest_id']==unknown[0]['id'] for u in job['result']['unassigned'])
    with SessionLocal() as db:assert not validate(snapshot(db),proposal,before=before)['errors']
    p=client.post('/api/planning/allocation/'+run_id+'/preview',json={'revision':revision(client)},headers=headers(owner)).json();assert p['token']
    assert apply(client,owner,p).status_code==200
    assert client.get('/api/planning/allocation').json()['runs'][0]['status']=='applied'
    assert client.post('/api/planning/allocation/'+run_id+'/preview',json={'revision':revision(client)},headers=headers(owner)).status_code==409


def test_worker_recovery_bounds_retries(client,owner):
    reserve(client,owner);family(client,owner)
    run_id=client.post('/api/planning/allocation',json={'revision':revision(client)},headers=headers(owner)).json()['id']
    with SessionLocal.begin() as db:
        run=db.get(AllocationRun,run_id);run.status='running';run.attempts=2;run.lease_until=now()-timedelta(seconds=1)
    process_one();assert client.get('/api/planning/allocation').json()['runs'][0]['status']=='failed'


def test_money_nights_staff_meals_ages_and_unknown_prices():
    guests=[{'id':str(i),'age':30,'arrival':NIGHTS[0],'category':'guest'} for i in range(132)]
    data={'guests':guests,'resources':[{'id':'w','hotel_id':'wyndham'},{'id':'s','hotel_id':'second'}],'assignments':[{'guest_id':str(i),'resource_id':'w','night':n} for n,number in [(NIGHTS[0],132),(NIGHTS[1],130)] for i in range(number)]}
    policy={'wyndham_allowance':128,'wyndham_excess_paise':500000,'staff_meal_paise':15000,'second_hotel_day_paise':1800000,'second_hotel_included':12};settings={'guest_meals':{},'second_extra_paise':None}
    result=calculate(data,policy,settings,[],[]);assert sum(l['amount_paise'] for l in result['lines'] if l['kind']=='wyndham_excess')==3000000
    data['guests'][0]['age']=10;data['guests'][1]['age']=11
    assert calculate(data,policy,settings,[],[])['nights'][NIGHTS[0]]['chargeable']==131
    staff=[{'id':'staff'+str(i),'age':30,'arrival':None,'category':'staff'} for i in range(5)];data['guests']+=staff
    meals=[{'guest_id':g['id'],'day':NIGHTS[0],'meal':meal,'location':'wyndham','planned':True,'actual':None,'inclusion':'chargeable'} for g in staff for meal in ['lunch','dinner']]
    priced=calculate(data,policy,settings,meals,[])
    assert sum(l['amount_paise'] for l in priced['lines'] if l['kind']=='staff_meal')==150000
    actual=calculate(data,policy,settings,meals,[],basis='actual');assert not any(l['kind']=='staff_meal' for l in actual['lines']) and any(u['kind']=='attendance_unknown' for u in actual['unpriced'])
    meals.append({'guest_id':'1','day':NIGHTS[0],'meal':'breakfast','location':'wyndham','planned':True,'actual':True,'inclusion':'included'})
    meals.append({'guest_id':'1','day':NIGHTS[0],'meal':'dinner','location':'wyndham','planned':True,'actual':True,'inclusion':'chargeable'})
    assert any(u['kind']=='guest_meal' for u in calculate(data,policy,settings,meals,[])['unpriced'])
    for i in range(13):data['assignments'][i]['resource_id']='s'
    hotel=calculate(data,policy,settings,meals,list(NIGHTS))
    assert sum(l['amount_paise'] for l in hotel['lines'] if l['kind']=='second_hotel_property')==3600000
    assert any(u['kind']=='second_extra' and u['quantity']==1 for u in hotel['unpriced']) and not hotel['complete']


def test_rates_snapshots_attendance_and_append_only_history(client,owner):
    guests=family(client,owner,[member('Synthetic Driver',30,category='driver')])
    meals={'revision':revision(client),'rows':[{'guest_id':guests[0]['id'],'day':NIGHTS[0],'meal':'dinner','planned':True,'actual':True,'inclusion':'chargeable'}]}
    assert client.put('/api/planning/meals',json=meals,headers=headers(owner)).status_code==200
    bill=client.get('/api/planning/billing').json();assert bill['known_subtotal_paise']==15000
    assert client.post('/api/planning/billing/snapshots',json={'revision':revision(client),'basis':'planned'},headers=headers(owner)).status_code==200
    rates={'revision':revision(client),**bill['policy'],**bill['settings'],'charge_days':[]};rates['staff_meal_paise']=16000
    assert client.put('/api/planning/billing/rates',json=rates,headers=headers(owner)).status_code==200
    assert client.get('/api/planning/billing').json()['known_subtotal_paise']==16000
    assert client.get('/api/planning/billing/snapshots').json()['snapshots'][0]['data']['known_subtotal_paise']==15000
    for sql in ['DELETE FROM bill_snapshots','DELETE FROM plan_changes']:
        with SessionLocal() as db:
            with pytest.raises(DBAPIError):db.execute(text(sql))


def workbook_bytes(rows):
    book=Workbook();sheet=book.active;sheet.title='Guests';sheet.append(HEADERS)
    for row in rows:sheet.append(row)
    out=io.BytesIO();book.save(out);return out.getvalue()


def test_excel_preview_apply_reimport_formula_safety_and_staleness(client,owner):
    body=workbook_bytes([[None,None,'Synthetic Imported','bride',False,False,'Synthetic Import Guest',35,'guest','+910000000000',NIGHTS[0],'second_hotel','Synthetic note']])
    def import_preview(content):return client.post('/api/planning/import/preview?revision='+str(revision(client)),files={'file':('guests.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},headers=headers(owner))
    result=import_preview(body);assert result.status_code==200,result.text;p=result.json();assert p['token'] and not p['errors'];assert not plan(client)['guests']
    r=client.post('/api/planning/import/apply',json={'revision':p['revision'],'token':p['token'],'acknowledged':[]},headers=headers(owner));assert r.status_code==200,r.text
    ids={g['id'] for g in plan(client)['guests']}
    exported=client.get('/api/planning/export.xlsx').content
    p=import_preview(exported).json();assert not p['errors'],p
    assert client.post('/api/planning/import/apply',json={'revision':p['revision'],'token':p['token'],'acknowledged':[w['id'] for w in p['warnings']]},headers=headers(owner)).status_code==200
    assert {g['id'] for g in plan(client)['guests']}==ids and plan(client)['guests'][0]['phone']=='+910000000000'
    p=import_preview(body).json();assert p['warnings']
    family(client,owner,label='Invalidate import')
    assert client.post('/api/planning/import/apply',json={'revision':p['revision'],'token':p['token'],'acknowledged':[w['id'] for w in p['warnings']]},headers=headers(owner)).status_code==409
    formula=workbook_bytes([[None,None,'Synthetic Formula','bride',False,False,'=1+1',30,'guest',None,NIGHTS[0],None,None]])
    assert import_preview(formula).json()['errors']
    person=family(client,owner,[member('=Synthetic text',30)],label='Export text')
    book=load_workbook(io.BytesIO(client.get('/api/planning/export.xlsx').content));cells=[c for row in book['Guests'] for c in row if c.value=="'=Synthetic text"]
    assert len(cells)==1 and cells[0].data_type!='f'


def test_migration_preserves_owner_and_session(client,owner):
    env={**os.environ,'PLANNER_DATABASE_USER':'planner_migrate','PLANNER_DATABASE_PASSWORD_FILE':'/run/secrets/db_migration_password'}
    for target in [('downgrade','0001_foundation'),('upgrade','head')]:
        result=subprocess.run(['alembic',*target],env=env,capture_output=True,text=True,timeout=30)
        assert result.returncode==0,result.stderr
    response=client.get('/api/auth/me');assert response.status_code==200 and response.json()['user']['id']==owner['user']['id']
    assert plan(client)['revision']==1 and not plan(client)['guests']


def test_live_revision_events_contain_no_guest_payload_and_expire(client,owner):
    family(client,owner,[member('PRIVATE SYNTHETIC PERSON',30,phone='PRIVATE CONTACT')])
    raw=client.cookies.get(COOKIE)
    async def disconnected():return False
    request=SimpleNamespace(cookies={COOKIE:raw},is_disconnected=disconnected)
    with SessionLocal() as db:user=db.get(User,owner['user']['id'])
    async def check():
        response=await events(request,user);stream=response.body_iterator
        first=await anext(stream);assert 'revision' in first and 'PRIVATE' not in first and 'paise' not in first
        with SessionLocal.begin() as db:
            row=db.get(LoginSession,digest(raw));db.delete(row)
        assert 'expired' in await anext(stream)
        await stream.aclose()
    asyncio.run(check())

def test_actual_arrivals_functions_and_side_scoped_exports(client,owner):
    bride=family(client,owner,[member('Synthetic Arriving Guest')])[0]
    groom=family(client,owner,[member('Synthetic Other Guest')],side='groom')[0]
    h=headers(owner)
    def record(g,arrival,departure=None):return client.put('/api/planning/presence/'+g['id'],json={'revision':revision(client),'arrived_at':arrival,'departed_at':departure},headers=h)
    assert record(bride,'2026-11-24T14:05:00+05:30').status_code==200
    assert record(bride,'2026-11-24T14:05:00+05:30','2026-11-24T12:00:00+05:30').status_code==400
    assert record(groom,'2026-11-25T14:00:00+05:30').status_code==200
    def create(name,side):
        r=client.post('/api/planning/events',json={'revision':revision(client),'name':name,'day':'2026-11-25','side':side},headers=h);assert r.status_code==201,r.text
        return r.json()['id']
    event=create('Synthetic Shared Reception',None);other=create('Synthetic Groom Function','groom')
    attendance={'event_id':event,'guest_id':bride['id'],'planned':True,'actual':True}
    assert client.put('/api/planning/functions/attendance',json={'revision':revision(client),'rows':[attendance]},headers=h).status_code==200
    bad={**attendance,'event_id':other}
    assert client.put('/api/planning/functions/attendance',json={'revision':revision(client),'rows':[bad]},headers=h).status_code==400
    assert client.post('/api/users',json={'username':'synthetic.bride','display_name':'Synthetic Bride Viewer','password':PASSWORD,'role':'family','side':'bride'},headers=h).status_code==201
    with TestClient(app) as viewer:
        auth=viewer.post('/api/auth/login',json={'username':'synthetic.bride','password':PASSWORD}).json()
        response=viewer.get('/api/planning');d=response.json()
        assert d['presence']==[{'guest_id':bride['id'],'status':'arrived'}]
        events=viewer.get('/api/planning/functions').json()
        assert {e['id'] for e in events['events']}=={event} and events['rows']==[attendance]
        assert viewer.put('/api/planning/presence/'+bride['id'],json={'revision':d['revision'],'arrived_at':None},headers={'X-CSRF-Token':auth['csrf_token']}).status_code==403
        book=load_workbook(io.BytesIO(viewer.get('/api/planning/export.xlsx').content))
        values=' '.join(str(c.value) for sheet in book for row in sheet for c in row)
        assert groom['id'] not in values and other not in values and '2026-11-24T08:35' not in values
        assert 'Arrivals' in book.sheetnames and 'Function attendance' in book.sheetnames


def test_one_meal_at_one_location_and_attendance_is_not_inferred(client,owner):
    g=family(client,owner,[member('Synthetic Meal Guest',arrival=None)])[0]
    row={'guest_id':g['id'],'day':'2026-11-25','meal':'dinner','location':'wyndham','planned':True,'actual':None,'inclusion':'unconfirmed'}
    assert client.put('/api/planning/meals',json={'revision':revision(client),'rows':[row]},headers=headers(owner)).status_code==200
    assert client.get('/api/planning/billing').json()['known_subtotal_paise']==0
    assert client.put('/api/planning/meals',json={'revision':revision(client),'rows':[{**row,'location':'external'}]},headers=headers(owner)).status_code==409
    meals=client.get('/api/planning/meals').json()['rows'];assert len(meals)==1 and meals[0]['actual'] is None
    # A single transaction can move a meal between locations without double counting.
    assert client.put('/api/planning/meals',json={'revision':revision(client),'rows':[{**row,'planned':False},{**row,'location':'external'}]},headers=headers(owner)).status_code==200

def test_known_age_bands_enable_provisional_assignment_without_invented_ages(client,owner):
    people=family(client,owner,[member(None,None,age_band='adult'),member(None,None,age_band='adult'),member(None,None,age_band='child')],couple_children=True)
    p=preview(client,owner,people,'w-102');assert not p['errors'] and not p['warnings']
    assert apply(client,owner,p).status_code==200
    d=plan(client);assert all(g['age'] is None for g in d['guests'])
    assert d['summary'][NIGHTS[0]]['unconfirmed_ages']==3
    assert len(d['conflicts']['notices'])==3
    assert client.get('/api/planning/billing').json()['nights'][NIGHTS[0]]['chargeable']==2
    # Confirming an elder's actual age cannot silently leave them upstairs.
    other=family(client,owner,[member(None,None,age_band='adult')],label='Synthetic Age Check')
    assign(client,owner,other,'w-202')
    f=next(f for f in plan(client)['families'] if f['id']==other[0]['family_id'])
    body={**f,'revision':revision(client),'members':[{**member(None,50,age_band='adult'),'id':other[0]['id']}]};body.pop('id')
    assert client.put('/api/planning/families/'+f['id'],json=body,headers=headers(owner)).status_code==409
    bad={**body,'revision':revision(client),'members':[{**member(None,10,age_band='adult'),'id':other[0]['id']}]}
    assert client.put('/api/planning/families/'+f['id'],json=bad,headers=headers(owner)).status_code==422


def test_unnamed_reserve_pool_is_shared_across_nights_and_solver_preserves_four(client,owner):
    with SessionLocal() as db:data=snapshot(db)
    premium=[r for r in data['resources'] if r['hotel_id']=='wyndham' and r['category']=='premium']
    data['resources']=premium;data['families']=[];data['guests']=[]
    for i in range(21):
        fid=f'synthetic-family-{i}';gid=f'synthetic-person-{i}'
        data['families'].append({'id':fid,'label':fid,'side':'bride','reserved_eligible':False,'couple_children':False})
        data['guests'].append({'id':gid,'family_id':fid,'name':None,'age':None,'age_band':'adult','side':'bride','reserved_eligible':False,'couple_children':False,'category':'guest','arrival':NIGHTS[0],'tags':[]})
    rows=[{'guest_id':g['id'],'night':NIGHTS[i%2],'resource_id':premium[i]['id']} for i,g in enumerate(data['guests'])]
    assert any(e['code']=='reserve_pool' for e in validate(data,rows)['errors'])
    result=allocate(data);assert result['assignments'] is not None,result
    assert len({a['resource_id'] for a in result['assignments']})<=20
    assert len(result['unassigned'])==2 and not validate(data,result['assignments'])['errors']
    assert sum(r['id'] not in {a['resource_id'] for a in result['assignments']} for r in premium)>=4


def test_automatic_pool_allocation_applies_and_is_not_reported_stale(client,owner):
    family(client,owner,[member(None,None,age_band='adult')])
    job=client.post('/api/planning/allocation',json={'revision':revision(client)},headers=headers(owner));assert job.status_code==202
    assert process_one()
    p=client.post('/api/planning/allocation/'+job.json()['id']+'/preview',json={'revision':revision(client)},headers=headers(owner)).json()
    assert p['token'] and apply(client,owner,p).status_code==200
    assert client.get('/api/planning/allocation').json()['runs'][0]['status']=='applied'
    exported=load_workbook(io.BytesIO(client.get('/api/planning/export.xlsx').content))
    assert 'Age band' in [c.value for c in exported['Guests'][1]]
    assert all(g['age'] is None and g['age_band']=='adult' for g in plan(client)['guests'])


def test_invitation_context_is_private_scoped_and_does_not_create_guests(client,owner):
    body={'revision':revision(client),'bride_name':'Synthetic Bride','groom_name':'Synthetic Groom','source_url':'https://example.com/invitation','family':[{'side':'bride','name':'Synthetic Bride Parent','relationship':'Father'},{'side':'groom','name':'Synthetic Groom Parent','relationship':'Mother'}]}
    assert client.put('/api/planning/wedding-details',json=body,headers=headers(owner)).status_code==200
    assert not plan(client)['guests']
    assert client.put('/api/planning/wedding-details',json=body,headers=headers(owner)).status_code==409
    invalid={**body,'revision':revision(client),'source_url':'javascript:alert(1)'}
    assert client.put('/api/planning/wedding-details',json=invalid,headers=headers(owner)).status_code==422
    for role,side in [('family','bride'),('reception',None)]:
        response=client.post('/api/users',json={'username':'context.'+role,'display_name':'Synthetic viewer','password':PASSWORD,'role':role,'side':side},headers=headers(owner));assert response.status_code==201
        with TestClient(app) as viewer:
            login=viewer.post('/api/auth/login',json={'username':'context.'+role,'password':PASSWORD});assert login.status_code==200
            details=viewer.get('/api/planning/wedding-details').json()
            assert len(details['family'])==(1 if role=='family' else 2)
            assert viewer.put('/api/planning/wedding-details',json={**body,'revision':revision(client)},headers=headers(login.json())).status_code==403
            assert viewer.get('/api/planning/billing').status_code==403


def test_function_times_sources_duplicate_guard_and_excel(client,owner):
    body={'revision':revision(client),'name':'Synthetic Midnight Ceremony','day':'2026-11-26','start_time':'00:00','venue':'Synthetic venue','source_url':'https://example.com/invitation'}
    response=client.post('/api/planning/events',json=body,headers=headers(owner));assert response.status_code==201,response.text
    body['revision']=revision(client)
    assert client.post('/api/planning/events',json=body,headers=headers(owner)).status_code==409
    assert client.post('/api/planning/events',json={**body,'start_time':'24:00'},headers=headers(owner)).status_code==422
    event=client.get('/api/planning/functions').json()['events'][0]
    assert event['day']=='2026-11-26' and event['start_time']=='00:00'
    wb=load_workbook(io.BytesIO(client.get('/api/planning/export.xlsx').content))
    assert wb['Functions']['E2'].value=='00:00' and wb['Functions']['F2'].value=='Synthetic venue'


def test_excel_optional_or_stale_dimensions_remain_bounded(client,owner):
    import re,zipfile
    from app.spreadsheet_api import parse_import
    def content(rows=1,far_column=False,dimension=None):
        wb=Workbook();sheet=wb.active;sheet.title='Guests';sheet.append(HEADERS)
        for i in range(rows):sheet.append([None,None,'Synthetic Import','bride',False,False,'Synthetic '+str(i),30,'guest',None,NIGHTS[0],'','', 'adult'])
        if far_column:sheet['AA2']='Synthetic excess column'
        out=io.BytesIO();wb.save(out);changed=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(out.getvalue())) as source,zipfile.ZipFile(changed,'w',zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data=source.read(item.filename)
                if item.filename=='xl/worksheets/sheet1.xml':data=re.sub(rb'<dimension[^>]*/>',dimension or b'',data)
                target.writestr(item,data)
        return changed.getvalue()
    current=plan(client)
    payload,errors,warnings=parse_import(content(),current);assert not errors and not warnings and len(payload['guests'])==1
    for invalid in [content(1001),content(far_column=True),content(1001,dimension=b'<dimension ref="A1:A1"/>')]:
        with pytest.raises(Exception) as caught:parse_import(invalid,current)
        assert caught.value.status_code==400 and '1,000' in caught.value.detail
